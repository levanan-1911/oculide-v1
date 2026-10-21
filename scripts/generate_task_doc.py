import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import os

def set_cell_background(cell, fill_hex):
    """Set background color for a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tc_pr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set cell margins in dxa (1 pt = 20 dxa)."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tc_mar.append(node)
    tc_pr.append(tc_mar)

def create_task_document(output_path):
    doc = Document()

    # 1. Cấu hình lề chuẩn Nghị định 30/2020/NĐ-CP (A4)
    # Lề trên: 2.0 cm, Lề dưới: 2.0 cm, Lề trái: 3.0 cm, Lề phải: 2.0 cm
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(3.0)
        section.right_margin = Cm(2.0)

    # 2. Định dạng Style mặc định: Times New Roman, 13pt
    normal_style = doc.styles['Normal']
    normal_font = normal_style.font
    normal_font.name = 'Times New Roman'
    normal_font.size = Pt(13)
    normal_font.color.rgb = RGBColor(30, 41, 59) # Slate 800

    # Header / Quốc hiệu hoặc Tên trường
    p_header = doc.add_paragraph()
    p_header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p_header.add_run("TRƯỜNG ĐẠI HỌC DUY TÂN - VIỆN ĐÀO TẠO QUỐC TẾ (CMU)\n")
    r1.font.bold = True
    r1.font.size = Pt(12)
    r1.font.color.rgb = RGBColor(71, 85, 105)
    
    r2 = p_header.add_run("ĐỒ ÁN TỐT NGHIỆP CAPSTONE 1 (C1SE.16) — HỌC KỲ I (2026-2027)\n")
    r2.font.bold = True
    r2.font.size = Pt(13)
    r2.font.color.rgb = RGBColor(15, 23, 42)

    r_line = p_header.add_run("———————————————\n")
    r_line.font.size = Pt(10)
    r_line.font.color.rgb = RGBColor(148, 163, 184)

    # Tiêu đề tài liệu
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(6)
    p_title.paragraph_format.space_after = Pt(4)
    r_t1 = p_title.add_run("KẾ HOẠCH PHÂN CHIA NHIỆM VỤ PHÁT TRIỂN DỰ ÁN\n")
    r_t1.font.bold = True
    r_t1.font.size = Pt(16)
    r_t1.font.color.rgb = RGBColor(15, 76, 129) # Classic Navy

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(14)
    r_s1 = p_sub.add_run("Đề tài: ")
    r_s1.font.italic = True
    r_s1.font.size = Pt(13)
    r_s2 = p_sub.add_run("Oculide: Ứng dụng thị giác máy tính và học sâu trong giám sát gian lận phòng thi lập trình trực tuyến\n")
    r_s2.font.bold = True
    r_s2.font.size = Pt(13)
    r_s2.font.color.rgb = RGBColor(2, 132, 199) # Blue
    r_s3 = p_sub.add_run("(Oculide: Application of Computer Vision and Deep Learning in Online Programming Exam Proctoring)")
    r_s3.font.italic = True
    r_s3.font.size = Pt(11)
    r_s3.font.color.rgb = RGBColor(100, 116, 139)

    # I. TỔNG QUAN PHÂN VAI & MA TRẬN NHIỆM VỤ
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(6)
    r_h1 = h1.add_run("I. TỔNG QUAN PHÂN BỔ NHIỆM VỤ NHÓM")
    r_h1.font.bold = True
    r_h1.font.size = Pt(14)
    r_h1.font.color.rgb = RGBColor(15, 76, 129)

    p_desc = doc.add_paragraph()
    p_desc.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_desc.paragraph_format.line_spacing = 1.2
    p_desc.paragraph_format.space_after = Pt(6)
    p_desc.add_run(
        "Hệ thống Oculide v1 được phát triển theo mô hình phân tán tích hợp (Multi-Tiered Architecture) gồm 3 phân hệ chính: "
        "Backend Core API & Docker Sandbox Auto-Grader, Động cơ AI Proctoring thị giác máy tính, và Frontend Web App (Next.js 14). "
        "Để đảm bảo tiến độ đồ án Capstone 1 và phát huy tối đa năng lực của từng thành viên, nhóm thống nhất bảng ma trận vai trò như sau:"
    )

    # Bảng 1: Ma trận phân vai
    table1 = doc.add_table(rows=6, cols=5)
    table1.alignment = WD_TABLE_ALIGNMENT.CENTER
    table1.autofit = False

    col_widths = [Cm(1.2), Cm(3.8), Cm(3.2), Cm(4.8), Cm(3.0)]
    headers = ["STT", "Họ và tên", "Vai trò chính", "Phân hệ phụ trách chính", "Nhánh Git chuẩn"]
    
    # Header row
    hdr_cells = table1.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        hdr_cells[i].paragraphs[0].runs[0].font.bold = True
        hdr_cells[i].paragraphs[0].runs[0].font.name = 'Times New Roman'
        hdr_cells[i].paragraphs[0].runs[0].font.size = Pt(12)
        hdr_cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        hdr_cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        hdr_cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_background(hdr_cells[i], "0F4C81") # Navy
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=100, right=100)

    # Data rows
    members_data = [
        ("1", "Lê Văn An", "AI, BE (Tech Lead)", "Core AI Engine, Docker Sandbox, Celery Queue & System Optimization", "feat/ai-core"),
        ("2", "Trần Thị Yến Nhi", "FE, BE, TESTING (Scrum Master)", "Student Exam Workspace, Monaco Online IDE & Submissions API", "feat/fe-exam-monaco"),
        ("3", "Huỳnh Trần Ngọc Trang", "FE, TEST INPUT", "Instructor Portal, Ngân hàng đề thi & Bộ dữ liệu Testcases", "feat/fe-instructor-data"),
        ("4", "Nguyễn Đình Lộc", "FE, TESTING", "Proctor Dashboard, Panoptic Video Grid (LiveKit) & Real-time Alert", "feat/fe-proctor-grid"),
        ("5", "Phạm Văn Long", "FE, TESTING", "Authentication, Admin Dashboard & QA Test Plan / Matrix", "feat/fe-auth-admin-qa")
    ]

    for row_idx, data in enumerate(members_data, start=1):
        row_cells = table1.rows[row_idx].cells
        bg_color = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(data):
            row_cells[col_idx].text = text
            p = row_cells[col_idx].paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            if p.runs:
                p.runs[0].font.name = 'Times New Roman'
                p.runs[0].font.size = Pt(12)
                p.runs[0].font.color.rgb = RGBColor(30, 41, 59)
                if col_idx in (0, 4):
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                elif col_idx == 2:
                    p.runs[0].font.bold = True
            row_cells[col_idx].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_background(row_cells[col_idx], bg_color)
            set_cell_margins(row_cells[col_idx], top=100, bottom=100, left=120, right=120)

    # Set column widths
    for row in table1.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = width

    # II. CHI TIẾT NHIỆM VỤ TỪNG THÀNH VIÊN
    h2 = doc.add_paragraph()
    h2.paragraph_format.space_before = Pt(16)
    h2.paragraph_format.space_after = Pt(6)
    r_h2 = h2.add_run("II. CHI TIẾT NHIỆM VỤ VÀ ĐẦU RA CỦA TỪNG THÀNH VIÊN")
    r_h2.font.bold = True
    r_h2.font.size = Pt(14)
    r_h2.font.color.rgb = RGBColor(15, 76, 129)

    # 1. Lê Văn An
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    p_m1 = doc.add_paragraph()
    p_m1.paragraph_format.space_before = Pt(6)
    p_m1.paragraph_format.space_after = Pt(3)
    r_m1 = p_m1.add_run("1. Lê Văn An — Core AI Engine & Backend Architecture (AI, BE)")
    r_m1.font.bold = True
    r_m1.font.size = Pt(13)
    r_m1.font.color.rgb = RGBColor(2, 132, 199)

    doc.add_paragraph("• Thư mục mã nguồn liên quan: ai_workers/, backend/ai_workers/, backend/tasks/, sandbox/, scripts/").paragraph_format.line_spacing = 1.2
    
    tasks_an = [
        ("Nghiên cứu & Huấn luyện thực nghiệm AI (Colab Training): ", 
         "Sử dụng kịch bản scripts/train_exam_yolo.py và Google Colab huấn luyện chuyên biệt mô hình YOLOv8 trên tập dữ liệu vật thể phòng thi (điện thoại, sách/phao, tai nghe). Trích xuất trọng số best.pt, đo đạc các chỉ số khoa học (mAP@50, Precision, Recall, F1-Score) để đưa vào tài liệu đồ án."),
        ("Tối ưu hóa bộ lọc gian lận thời gian thực (AI Proctoring Core): ", 
         "Duy trì và tối ưu mã nguồn ai_workers/detector.py: MediaPipe FaceMesh & Iris Gaze Tracking (478 điểm mốc 3D), ước lượng tư thế đầu Pitch/Yaw/Roll qua SolvePnP, thuật toán EAR chống giả mạo, máy trạng thái chuỗi thời gian (quy tắc Dwell-time >= 7s) và cơ chế tương quan gõ phím - hướng nhìn (Keystroke-Gaze Correlation)."),
        ("Tối ưu hóa Docker Sandbox & Celery Auto-Grader: ", 
         "Cấu hình các container thực thi cô lập (Python 3.11, GCC 12, OpenJDK 17) trong backend/tasks/grading_tasks.py. Thiết lập giới hạn bộ nhớ (128MB), giới hạn CPU, ngắt kết nối internet (network_disabled: true) và cơ chế timeout 10s ngăn chặn mã độc/fork-bomb."),
        ("Quản lý kiến trúc hệ thống, WebSocket & Review Code: ", 
         "Đảm bảo hạ tầng kết nối thời gian thực qua WebSocket /ws/{room_id}/{user_id}, cấp phát LiveKit Token bảo mật. Đóng vai trò Tech Lead kiểm duyệt Pull Requests (PR) từ các thành viên vào nhánh main, giải quyết xung đột mã nguồn và hỗ trợ kỹ thuật.")
    ]
    for title, desc in tasks_an:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.line_spacing = 1.2
        p.paragraph_format.space_after = Pt(3)
        r_b = p.add_run(title)
        r_b.font.bold = True
        p.add_run(desc)

    # 2. Trần Thị Yến Nhi
    p_m2 = doc.add_paragraph()
    p_m2.paragraph_format.space_before = Pt(8)
    p_m2.paragraph_format.space_after = Pt(3)
    r_m2 = p_m2.add_run("2. Trần Thị Yến Nhi — Student Exam Workspace & Monaco IDE (FE, BE, TESTING)")
    r_m2.font.bold = True
    r_m2.font.size = Pt(13)
    r_m2.font.color.rgb = RGBColor(2, 132, 199)

    doc.add_paragraph("• Thư mục mã nguồn liên quan: frontend/src/app/exam/, backend/api/submissions.py, backend/api/sessions.py").paragraph_format.line_spacing = 1.2

    tasks_nhi = [
        ("Khởi tạo nền tảng Web Frontend Next.js 14: ", 
         "Setup cấu trúc chuẩn App Router, TypeScript, Tailwind CSS, Lucide-react, Zustand store và Axios interceptor tự động đính kèm JWT Token vào request header."),
        ("Xây dựng Giao diện phòng thi trực tuyến (/app/exam): ", 
         "Thiết kế bố cục chia đôi màn hình (Split-Pane): Cột trái hiển thị đề bài lập trình (Markdown rendering, danh sách ví dụ mẫu); Cột phải tích hợp trình soạn thảo Monaco Editor (chọn ngôn ngữ Python/C++/Java, đổi theme sáng/tối, phím tắt); Phía dưới là Console hiển thị kết quả biên dịch và trạng thái test cases."),
        ("Tích hợp nút Chạy thử (Run Code) & Nộp bài (Submit): ", 
         "Gọi API /api/v1/submissions/run để chạy thử code với custom input mà không tính điểm; Gọi API /api/v1/submissions để nộp bài chính thức, nhận trạng thái chấm điểm từ Celery Worker."),
        ("Kiểm soát hành vi gian lận trình duyệt & Chụp webcam ngầm: ", 
         "Bắt các sự kiện gian lận: chuyển tab (visibilitychange), mất tiêu điểm (blur), thoát toàn màn hình (fullscreenchange), chặn chuột phải và chặn phím tắt copy/paste phao code. Nhúng webcam ngầm gọi module proctorClient.ts định kỳ nén ảnh 480p gửi về backend phân tích."),
        ("Backend & Testing: ", 
         "Phối hợp hoàn thiện API Submissions, viết unit test kiểm thử tính toàn vẹn của mã nguồn nộp và tính toán tổng điểm bài thi.")
    ]
    for title, desc in tasks_nhi:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.line_spacing = 1.2
        p.paragraph_format.space_after = Pt(3)
        r_b = p.add_run(title)
        r_b.font.bold = True
        p.add_run(desc)

    # 3. Huỳnh Trần Ngọc Trang
    p_m3 = doc.add_paragraph()
    p_m3.paragraph_format.space_before = Pt(8)
    p_m3.paragraph_format.space_after = Pt(3)
    r_m3 = p_m3.add_run("3. Huỳnh Trần Ngọc Trang — Instructor Portal & Test Input Data (FE, TEST INPUT)")
    r_m3.font.bold = True
    r_m3.font.size = Pt(13)
    r_m3.font.color.rgb = RGBColor(2, 132, 199)

    doc.add_paragraph("• Thư mục mã nguồn liên quan: frontend/src/app/instructor/, database/seeds/, docs/problem_bank/").paragraph_format.line_spacing = 1.2

    tasks_trang = [
        ("Giao diện Quản lý phòng thi & Lớp học (/app/instructor/room): ", 
         "Xây dựng trang tạo phòng, chọn chế độ (Luyện tập classroom hoặc Thi cử exam), cấu hình thời gian mở/đóng phòng, thời lượng làm bài, passcode, quản lý danh sách sinh viên ghi danh (Enrollments)."),
        ("Giao diện Soạn thảo đề bài & Ngân hàng câu hỏi (/app/instructor/questions): ", 
         "Xây dựng form tạo câu hỏi hỗ trợ định dạng Markdown, cấu hình giới hạn thời gian chạy (Time Limit) và bộ nhớ (Memory Limit) cho từng bài toán."),
        ("Xây dựng Bảng điểm thời gian thực (Scoreboard / Leaderboard): ", 
         "Bảng tổng hợp điểm số trực tiếp của toàn bộ sinh viên trong phòng thi: số câu giải được, điểm số, thời gian nộp bài, số lần vi phạm. Tích hợp tính năng Xuất báo cáo (Export Excel/CSV)."),
        ("TEST INPUT — Xây dựng Ngân hàng Đề thi & Bộ dữ liệu kiểm thử (Test Cases Suite): ", 
         "Thiết kế 15-20 bài toán lập trình chuẩn (nhập xuất, mảng, xâu, đệ quy, quy hoạch động). Với mỗi bài toán, xây dựng trọn vẹn: 2-3 Sample Test Cases (công khai cho thí sinh chạy thử); 5-10 Hidden Test Cases (bí mật để chấm điểm); các Edge Cases (dữ liệu rỗng, số âm, giá trị cực đại) và Big Data Cases (kiểm tra thuật toán tối ưu, chặn lỗi TLE/MLE). Viết file script nạp sẵn dữ liệu seed vào CSDL SQL Server.")
    ]
    for title, desc in tasks_trang:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.line_spacing = 1.2
        p.paragraph_format.space_after = Pt(3)
        r_b = p.add_run(title)
        r_b.font.bold = True
        p.add_run(desc)

    # 4. Nguyễn Đình Lộc
    p_m4 = doc.add_paragraph()
    p_m4.paragraph_format.space_before = Pt(8)
    p_m4.paragraph_format.space_after = Pt(3)
    r_m4 = p_m4.add_run("4. Nguyễn Đình Lộc — Proctoring Dashboard & Panoptic Grid (FE, TESTING)")
    r_m4.font.bold = True
    r_m4.font.size = Pt(13)
    r_m4.font.color.rgb = RGBColor(2, 132, 199)

    doc.add_paragraph("• Thư mục mã nguồn liên quan: frontend/src/app/proctor/, frontend/src/hooks/useLiveKit.ts, useWebSocket.ts").paragraph_format.line_spacing = 1.2

    tasks_loc = [
        ("Xây dựng Panoptic Grid View giám sát phòng thi (/app/proctor): ", 
         "Ứng dụng LiveKit WebRTC SDK (@livekit/components-react) để dựng lưới video webcam hiển thị trực tiếp tất cả thí sinh trong phòng thi theo thời gian thực (Grid 3x3 hoặc 4x4 có phân trang). Hiển thị trạng thái online/offline và huy hiệu trạng thái của từng sinh viên."),
        ("Real-time Violation Alert Hub: ", 
         "Kết nối WebSocket /ws/{room_id}/{user_id} nhận thông báo cảnh báo tức thì từ AI Proctoring khi có sinh viên quay đầu, liếc mắt >= 7s, mất khuôn mặt hoặc xuất hiện điện thoại/sách phao."),
        ("Modal xem xét bằng chứng vi phạm (Violation Evidence Review): ", 
         "Khi giám thị nhấp vào cảnh báo vi phạm, hiển thị popup phóng to ảnh snapshot webcam lúc vi phạm kèm thông tin góc lệch và thời gian duy trì. Cung cấp nút duyệt: 'Xác nhận vi phạm' hoặc 'Bỏ qua cảnh báo' gọi API /api/v1/violations/{id}/review."),
        ("Khung chatbox phòng thi & Kiểm thử hiệu năng (Testing): ", 
         "Tích hợp chatbox trao đổi giữa giám thị và thí sinh trong phòng. Thực hiện kiểm thử tải (Stress Testing) truyền phát nhiều luồng video WebRTC và đo đạc độ trễ nhận thông báo cảnh báo qua WebSocket (yêu cầu độ trễ < 2s).")
    ]
    for title, desc in tasks_loc:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.line_spacing = 1.2
        p.paragraph_format.space_after = Pt(3)
        r_b = p.add_run(title)
        r_b.font.bold = True
        p.add_run(desc)

    # 5. Phạm Văn Long
    p_m5 = doc.add_paragraph()
    p_m5.paragraph_format.space_before = Pt(8)
    p_m5.paragraph_format.space_after = Pt(3)
    r_m5 = p_m5.add_run("5. Phạm Văn Long — Authentication, Admin Dashboard & QA/Testing (FE, TESTING)")
    r_m5.font.bold = True
    r_m5.font.size = Pt(13)
    r_m5.font.color.rgb = RGBColor(2, 132, 199)

    doc.add_paragraph("• Thư mục mã nguồn liên quan: frontend/src/app/login/, register/, profile/, join/, admin/, docs/testing/").paragraph_format.line_spacing = 1.2

    tasks_long = [
        ("Phân hệ Xác thực & Quản lý tài khoản cá nhân: ", 
         "Xây dựng trang Đăng nhập (/app/login), Đăng ký (/app/register), Quên mật khẩu, Trang thông tin cá nhân (/app/profile) cho phép cập nhật họ tên và đổi mật khẩu. Lưu trữ JWT Access Token an toàn vào HttpOnly Cookie / localStorage."),
        ("Trang tham gia nhanh bằng mã phòng (/app/join): ", 
         "Giao diện nhập mã phòng (room_code) và passcode để sinh viên vào thẳng phòng thi mà không cần tìm kiếm danh sách."),
        ("Dashboard Quản trị Hệ thống (/app/admin): ", 
         "Trang quản lý người dùng: xem danh sách có phân trang và tìm kiếm, lọc theo vai trò (admin, instructor, student), tạo tài khoản, khóa/mở tài khoản, reset mật khẩu, chức năng Batch Import danh sách sinh viên qua file CSV/Excel. Trang xem Nhật ký hệ thống (System Logs / Audit Logs)."),
        ("QA & Xây dựng Kế hoạch Kiểm thử tổng thể (Testing Documentation): ", 
         "Xây dựng tài liệu Kế hoạch kiểm thử (Test Plan) và Ma trận kịch bản kiểm thử (Test Matrix) bao phủ toàn bộ chức năng hệ thống: kiểm thử phân quyền RBAC (ngăn chặn sinh viên truy cập trái phép trang quản trị), kiểm thử Form validation, kiểm thử giao diện responsive trên nhiều độ phân giải màn hình. Ghi nhận lỗi và báo cáo bug lên GitHub Issues.")
    ]
    for title, desc in tasks_long:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.line_spacing = 1.2
        p.paragraph_format.space_after = Pt(3)
        r_b = p.add_run(title)
        r_b.font.bold = True
        p.add_run(desc)

    # III. QUY TRÌNH PHỐI HỢP VÀ NGUYÊN TẮC LÀM VIỆC NHÓM
    h3 = doc.add_paragraph()
    h3.paragraph_format.space_before = Pt(16)
    h3.paragraph_format.space_after = Pt(6)
    r_h3 = h3.add_run("III. QUY TRÌNH PHỐI HỢP VÀ NGUYÊN TẮC LÀM VIỆC NHÓM")
    r_h3.font.bold = True
    r_h3.font.size = Pt(14)
    r_h3.font.color.rgb = RGBColor(15, 76, 129)

    guidelines = [
        ("Quy tắc phân nhánh Git (Branching Strategy): ", 
         "Nhánh 'main' là nhánh chính thức được bảo vệ (Protected Branch). Các thành viên tạo nhánh tính năng từ 'main' theo định dạng: feat/<tên-tính-năng>. Tuyệt đối không commit trực tiếp vào 'main'."),
        ("Quy chuẩn thông điệp Commit (Conventional Commits): ", 
         "Sử dụng tiền tố chuẩn: feat: (tính năng mới), fix: (sửa lỗi), docs: (tài liệu), test: (kiểm thử), refactor: (tái cấu trúc code). Ví dụ: feat(exam): implement monaco editor split pane layout."),
        ("Quy trình Review và Gộp mã nguồn (Pull Request & Merge): ", 
         "Khi hoàn thành tính năng, thành viên đẩy nhánh lên GitHub và tạo Pull Request (PR) trỏ vào 'main'. Tech Lead (Lê Văn An) và Scrum Master (Trần Thị Yến Nhi) sẽ review code, kiểm tra CI automated test trước khi bấm Merge."),
        ("Họp Sprint định kỳ (Agile/Scrum): ", 
         "Tổ chức họp ngắn (Daily Standup / Weekly Review) vào đầu tuần để tổng kết công việc đã làm, xác định việc cần làm tiếp theo và cùng giải quyết các vướng mắc kỹ thuật (blockers).")
    ]
    for title, desc in guidelines:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.line_spacing = 1.2
        p.paragraph_format.space_after = Pt(3)
        r_b = p.add_run(title)
        r_b.font.bold = True
        p.add_run(desc)

    # IV. LỘ TRÌNH TRIỂN KHAI VÀ MỤC TIÊU CỘT MỐC (ROADMAP)
    h4 = doc.add_paragraph()
    h4.paragraph_format.space_before = Pt(16)
    h4.paragraph_format.space_after = Pt(6)
    r_h4 = h4.add_run("IV. LỘ TRÌNH THỰC HIỆN VÀ CỘT MỐC GIAO NỘP (ROADMAP)")
    r_h4.font.bold = True
    r_h4.font.size = Pt(14)
    r_h4.font.color.rgb = RGBColor(15, 76, 129)

    table2 = doc.add_table(rows=5, cols=4)
    table2.alignment = WD_TABLE_ALIGNMENT.CENTER
    table2.autofit = False

    t2_widths = [Cm(2.2), Cm(3.8), Cm(6.5), Cm(3.5)]
    t2_headers = ["Giai đoạn", "Mục tiêu trọng tâm", "Nội dung triển khai", "Sản phẩm bàn giao"]
    
    t2_hdr = table2.rows[0].cells
    for i, title in enumerate(t2_headers):
        t2_hdr[i].text = title
        t2_hdr[i].paragraphs[0].runs[0].font.bold = True
        t2_hdr[i].paragraphs[0].runs[0].font.name = 'Times New Roman'
        t2_hdr[i].paragraphs[0].runs[0].font.size = Pt(12)
        t2_hdr[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        t2_hdr[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        t2_hdr[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_background(t2_hdr[i], "0F4C81")
        set_cell_margins(t2_hdr[i], top=120, bottom=120, left=100, right=100)

    roadmap_data = [
        ("Tuần 1 - 2", "Khởi tạo & Kiến trúc lõi (Đã hoàn thành)", "Thiết kế CSDL 13 bảng, chuẩn hóa RESTful API v1, tích hợp động cơ AI MediaPipe + YOLOv8, Docker Sandbox.", "Bộ Backend, Docker Compose chạy 5 container healthy, 31 unit tests passed."),
        ("Tuần 3 - 5", "Xây dựng Giao diện Web & Monaco IDE (Hiện tại)", "Khởi tạo Next.js 14, hoàn thiện Auth, giao diện Exam Monaco IDE, Instructor Portal và bộ Testcases ban đầu.", "Màn hình đăng nhập, phòng thi Monaco IDE chạy thử code được, giao diện tạo phòng."),
        ("Tuần 6 - 8", "Tích hợp Streaming LiveKit & Panoptic Grid", "Ghép nối WebRTC LiveKit, hoàn thiện lưới webcam giám thị, nhận cảnh báo vi phạm WebSocket thời gian thực.", "Dashboard Giám thị hiển thị video trực tiếp và popup bằng chứng vi phạm AI."),
        ("Tuần 9 - 12", "Kiểm thử toàn diện, Đóng gói & Báo cáo", "Train hoàn tất YOLOv8 Colab, Stress test hệ thống, hoàn thiện hồ sơ kiểm thử QA, slide và video demo.", "Hệ thống Oculide v1 hoàn chỉnh, tài liệu báo cáo đồ án Capstone 1 và video demo.")
    ]

    for r_idx, r_data in enumerate(roadmap_data, start=1):
        r_cells = table2.rows[r_idx].cells
        bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(r_data):
            r_cells[c_idx].text = val
            p = r_cells[c_idx].paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            if p.runs:
                p.runs[0].font.name = 'Times New Roman'
                p.runs[0].font.size = Pt(12)
                p.runs[0].font.color.rgb = RGBColor(30, 41, 59)
                if c_idx == 0:
                    p.runs[0].font.bold = True
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r_cells[c_idx].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_background(r_cells[c_idx], bg_color)
            set_cell_margins(r_cells[c_idx], top=100, bottom=100, left=120, right=120)

    for row in table2.rows:
        for idx, width in enumerate(t2_widths):
            row.cells[idx].width = width

    # Footer chữ ký
    p_sign = doc.add_paragraph()
    p_sign.paragraph_format.space_before = Pt(24)
    p_sign.paragraph_format.line_spacing = 1.2
    
    t_sign = doc.add_table(rows=1, cols=2)
    t_sign.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_sign.autofit = False
    t_sign.rows[0].cells[0].width = Cm(8.0)
    t_sign.rows[0].cells[1].width = Cm(8.0)

    c_left = t_sign.rows[0].cells[0]
    p_l = c_left.paragraphs[0]
    p_l.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_l1 = p_l.add_run("XÁC NHẬN NHÓM TRƯỞNG\n")
    r_l1.font.bold = True
    r_l1.font.size = Pt(12)
    p_l.add_run("(Ký và ghi rõ họ tên)\n\n\n\n")
    r_l2 = p_l.add_run("Trần Thị Yến Nhi")
    r_l2.font.bold = True
    r_l2.font.size = Pt(12)

    c_right = t_sign.rows[0].cells[1]
    p_r = c_right.paragraphs[0]
    p_r.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_r1 = p_r.add_run("ĐẠI DIỆN KỸ THUẬT / TECH LEAD\n")
    r_r1.font.bold = True
    r_r1.font.size = Pt(12)
    p_r.add_run("(Ký và ghi rõ họ tên)\n\n\n\n")
    r_r2 = p_r.add_run("Lê Văn An")
    r_r2.font.bold = True
    r_r2.font.size = Pt(12)

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    print(f"Document successfully created at: {output_path}")

if __name__ == "__main__":
    out_file = r"d:\oculide-v1\docs\Ke_Hoach_Phan_Chia_Cong_Viec_Oculide_v1.docx"
    create_task_document(out_file)
