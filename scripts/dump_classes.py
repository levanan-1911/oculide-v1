import docx

doc = docx.Document(r'docs/C1SE.16_SoftwareRequirementSpecification_OCLIDE_VIE_ver1.0.docx')

with open(r'docs/domain_classes.txt', 'w', encoding='utf-8') as f:
    for idx in [33, 34, 35]:
        if idx < len(doc.tables):
            tbl = doc.tables[idx]
            f.write(f"\n==================== TABLE {idx} (rows: {len(tbl.rows)}, cols: {len(tbl.columns)}) ====================\n")
            for r_idx, row in enumerate(tbl.rows):
                row_txt = [c.text.strip().replace('\n', ' ') for c in row.cells]
                f.write(f"R{r_idx}: {' | '.join(row_txt)}\n")

print("Dumped domain classes successfully")
