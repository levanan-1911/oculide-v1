import docx

doc = docx.Document(r'docs/C1SE.16_SoftwareRequirementSpecification_OCLIDE_VIE_ver1.0.docx')
with open(r'docs/srs_content.txt', 'w', encoding='utf-8') as f:
    f.write(f'TOTAL PARAGRAPHS: {len(doc.paragraphs)}\n')
    f.write(f'TOTAL TABLES: {len(doc.tables)}\n\n')
    for i, p in enumerate(doc.paragraphs):
        txt = p.text.strip()
        if txt:
            f.write(f'P{i}: {txt}\n')
    for t_idx, table in enumerate(doc.tables):
        f.write(f'\n--- TABLE {t_idx} (rows: {len(table.rows)}, cols: {len(table.columns)}) ---\n')
        for r_idx, row in enumerate(table.rows):
            row_txt = [c.text.strip().replace('\n', ' ') for c in row.cells]
            f.write(f'  R{r_idx}: {" | ".join(row_txt)}\n')
print('Dumped successfully to docs/srs_content.txt')
